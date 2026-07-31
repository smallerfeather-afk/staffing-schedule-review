"""Tests for cleaned Excel workbook generation."""

from datetime import datetime, time
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pandas as pd

from staffing_review.constants import REVIEW_COLUMNS
from staffing_review.io.exporters import export_filename, export_reviewed_schedule


def _review_records() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "source_row": 2,
                "employee_name": "Avery Stone",
                "role": "Caregiver",
                "shift_date": "2026-08-03",
                "start_time": "07:00",
                "unit": "Assisted Living",
                "floor": "1",
                "validation_status": "Valid",
                "validation_messages": "",
            },
            {
                "source_row": 3,
                "employee_name": "OPEN",
                "role": "Med Tech",
                "shift_date": "2026-08-03",
                "start_time": "15:00",
                "unit": "Memory Care",
                "floor": "2A",
                "validation_status": "Valid",
                "validation_messages": "",
            },
            {
                "source_row": 4,
                "employee_name": "Morgan Reed",
                "role": "RN",
                "shift_date": "bad-date",
                "start_time": "25:00",
                "unit": "Skilled Nursing",
                "floor": "Ground",
                "validation_status": "Invalid",
                "validation_messages": "Shift date is invalid; Start time is invalid",
            },
        ]
    )


def test_export_reopens_with_typed_values_and_highlights() -> None:
    payload = export_reviewed_schedule(_review_records())

    workbook = load_workbook(BytesIO(payload))
    worksheet = workbook["Reviewed Schedule"]

    assert workbook.sheetnames == ["Reviewed Schedule"]
    assert [cell.value for cell in worksheet[1]] == list(REVIEW_COLUMNS)
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:I4"
    assert worksheet["D2"].value == datetime(2026, 8, 3)
    assert worksheet["E2"].value == time(7, 0)
    assert worksheet["H3"].value == "Open"
    assert worksheet["A3"].fill.fgColor.rgb == "00FFF1CC"
    assert worksheet["A4"].fill.fgColor.rgb == "00FCE8E6"
    assert worksheet["I4"].value == (
        "Shift date is invalid; Start time is invalid"
    )


def test_export_preserves_mapped_xlsx_styles() -> None:
    source = Workbook()
    worksheet = source.active
    worksheet.append(
        ["Employee", "Date", "Start Time", "Department", "Level"]
    )
    worksheet.append(
        ["Avery Stone", "2026-08-03", "07:00", "AL", "1"]
    )
    worksheet["A1"].fill = PatternFill("solid", fgColor="336699")
    worksheet["A2"].font = Font(name="Calibri", bold=True, color="123456")
    worksheet.column_dimensions["A"].width = 31
    worksheet.row_dimensions[2].height = 28
    source_bytes = BytesIO()
    source.save(source_bytes)

    payload = export_reviewed_schedule(
        _review_records().iloc[[0]],
        source_workbook=source_bytes.getvalue(),
    )
    exported = load_workbook(BytesIO(payload))["Reviewed Schedule"]

    assert exported["B1"].fill.fgColor.rgb == "00336699"
    assert exported["B2"].font.bold is True
    assert exported["B2"].font.color.rgb == "00123456"
    assert exported.column_dimensions["B"].width == 31
    assert exported.row_dimensions[2].height == 28


def test_export_treats_formula_like_names_as_text() -> None:
    records = _review_records().iloc[[0]].copy()
    records.loc[records.index[0], "employee_name"] = "=HYPERLINK(\"x\",\"x\")"

    payload = export_reviewed_schedule(records)
    worksheet = load_workbook(BytesIO(payload), data_only=False).active

    assert worksheet["B2"].data_type == "s"
    assert worksheet["B2"].value == "'=HYPERLINK(\"x\",\"x\")"


def test_export_filename_is_windows_safe() -> None:
    assert export_filename("August SNF schedule.xlsx") == (
        "August_SNF_schedule_reviewed.xlsx"
    )
    assert export_filename("<>") == "staffing_schedule_reviewed.xlsx"


def test_export_preserves_pdf_extraction_metadata() -> None:
    records = _review_records().iloc[[0]].copy()
    records["source_page"] = 4
    records["extraction_confidence"] = 0.75
    records["extraction_notes"] = "Location was not found"
    records["extraction_reviewed"] = True

    payload = export_reviewed_schedule(records)
    worksheet = load_workbook(BytesIO(payload)).active

    assert worksheet.max_column == 13
    assert [worksheet.cell(1, column).value for column in range(10, 14)] == [
        "source_page",
        "extraction_confidence",
        "extraction_notes",
        "extraction_reviewed",
    ]
    assert worksheet["J2"].value == 4
    assert worksheet["K2"].value == 0.75
    assert worksheet["L2"].value == "Location was not found"
    assert worksheet["M2"].value is True
