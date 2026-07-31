"""Tests for population of the approved SNF two-week template."""

from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from staffing_review.io.snf_template import (
    BUILT_IN_TEMPLATE_PATH,
    RN_FONT_COLOR,
    SNFTemplateError,
    export_snf_template,
    snf_template_filename,
)


def _records() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "employee_name": "James CNA",
                "role": "CNA",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Skilled Nursing",
                "floor": "S1 James",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Shared CNA",
                "role": "CNA",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Skilled Nursing",
                "floor": "S1 James- Johns",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Folly CNA A",
                "role": "CNA",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Folly",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Folly CNA B",
                "role": "CNA",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Folly",
                "validation_status": "Valid",
            },
            {
                "employee_name": "James LPN",
                "role": "LPN",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Skilled Nursing",
                "floor": "S1 James",
                "validation_status": "Valid",
            },
            {
                "employee_name": "James RN",
                "role": "RN",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Skilled Nursing",
                "floor": "S1 James",
                "validation_status": "Valid",
            },
            {
                "employee_name": "House RN",
                "role": "RN",
                "shift_date": "2026-07-31",
                "start_time": "08:00",
                "unit": "Skilled Nursing",
                "floor": "",
                "validation_status": "Invalid",
            },
            {
                "employee_name": "Week Two CNA",
                "role": "CNA",
                "shift_date": "2026-08-07",
                "start_time": "07:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Morris",
                "validation_status": "Valid",
            },
        ]
    )


def test_populates_cna_nurse_and_week_two_rows_excluding_invalid() -> None:
    output = export_snf_template(_records())
    worksheet = load_workbook(BytesIO(output)).active

    assert worksheet["B2"].value == "Friday, 7/31 - AM"
    assert worksheet["O22"].value == "Thursday, 8/13 - PM"
    assert worksheet["B3"].value == "James CNA"
    assert worksheet["B4"].value == "Shared CNA"
    assert worksheet["C4"].value == "XXXXXXXXXXXX"
    assert worksheet["C6"].value == "Folly CNA A\nFolly CNA B"
    assert worksheet["C6"].alignment.wrap_text is True
    assert worksheet["B18"].value == "James LPN\nJames RN"
    assert worksheet["B19"].value == "OPEN"
    assert worksheet["B20"].value == "OPEN"
    assert worksheet["B21"].value == "OPEN"
    assert worksheet["B15"].value == "Week Two CNA"


def test_applies_snf_font_and_dimension_rules() -> None:
    worksheet = load_workbook(
        BytesIO(export_snf_template(_records()))
    ).active

    assert worksheet["D3"].value == "OPEN"
    assert worksheet["D3"].font.bold is True
    assert worksheet["D3"].font.color.rgb == "FFFF0000"
    assert worksheet["B3"].font.bold is False
    assert worksheet["B3"].font.color.rgb == "FF000000"
    assert worksheet["C6"].font.bold is False
    assert worksheet["C6"].font.color.rgb == "FFFFFFFF"
    assert worksheet["C4"].font.color.rgb == "FFFFFF00"
    assert worksheet.row_dimensions[3].height == 48
    assert worksheet.column_dimensions["A"].width == 14
    assert worksheet.column_dimensions["B"].width == 22
    assert worksheet["A3"].font.sz == 12
    assert worksheet["B2"].font.sz == 10
    assert worksheet["B1"].font.sz == 24


def test_excludes_invalid_records_and_rejects_empty_valid_set() -> None:
    records = _records()
    records.loc[0, "floor"] = "Unknown Hall"
    records.loc[0, "validation_status"] = "Invalid"

    worksheet = load_workbook(
        BytesIO(export_snf_template(records))
    ).active
    assert worksheet["B3"].value == "OPEN"

    records["validation_status"] = "Invalid"
    with pytest.raises(SNFTemplateError, match="no valid reviewed records"):
        export_snf_template(records)


def test_merges_open_folly_pm_with_named_morris_and_colors_rn() -> None:
    records = pd.DataFrame(
        [
            {
                "employee_name": "OPEN",
                "role": "LPN",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Folly",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Morris LPN",
                "role": "LPN",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Morris",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Morris RN",
                "role": "RN",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Skilled Nursing",
                "floor": "S2 Morris",
                "validation_status": "Valid",
            },
        ]
    )

    worksheet = load_workbook(
        BytesIO(export_snf_template(records)),
        rich_text=True,
    ).active
    assert "C20:C21" in {str(item) for item in worksheet.merged_cells.ranges}
    assert str(worksheet["C20"].value) == "Morris LPN\nMorris RN"
    assert "OPEN" not in str(worksheet["C20"].value)
    assert "(RN)" not in str(worksheet["C20"].value)
    rn_blocks = [
        block
        for block in worksheet["C20"].value
        if getattr(block, "text", "") == "Morris RN"
    ]
    assert len(rn_blocks) == 1
    assert rn_blocks[0].font.color.rgb == RN_FONT_COLOR


def test_built_in_template_and_filename() -> None:
    assert BUILT_IN_TEMPLATE_PATH.is_file()
    assert snf_template_filename("SNF weekly report.pdf") == (
        "SNF_weekly_report_SNF_template.xlsx"
    )
