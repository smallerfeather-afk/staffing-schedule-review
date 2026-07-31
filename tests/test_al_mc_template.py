"""Tests for population of the approved AL/MC two-week template."""

from copy import copy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pandas as pd
import pytest

from staffing_review.io.al_mc_template import (
    ALMCTemplateError,
    BUILT_IN_TEMPLATE_PATH,
    al_mc_template_filename,
    export_al_mc_template,
)


def _template_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.merge_cells("B1:O1")
    worksheet["B1"] = "Assisted Living"
    worksheet.merge_cells("B23:O23")
    worksheet["B23"] = "Memory Care"

    sections = [
        (2, ["AL1", "AL2", "AL3"], {"AL1": 2, "AL2": 2, "AL3": 2}),
        (12, ["AL1", "AL2", "AL3"], {"AL1": 2, "AL2": 2, "AL3": 2}),
        (24, ["MC1", "MC2"], {"MC1": 3, "MC2": 2}),
        (32, ["MC1", "MC2"], {"MC1": 3, "MC2": 2}),
    ]
    for header_row, floors, slots_by_floor in sections:
        worksheet.cell(header_row, 1, "Hall")
        row = header_row + 1
        for floor in floors:
            worksheet.cell(row, 1, f"{floor} - MT")
            row += 1
            for slot in range(1, slots_by_floor[floor] + 1):
                worksheet.cell(row, 1, f"{floor} - CG{slot}")
                row += 1

    for row in range(1, 40):
        for column in range(2, 16):
            cell = worksheet.cell(row, column)
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            cell.alignment = alignment
            if column % 2:
                cell.fill = PatternFill("solid", fgColor="000000")
                cell.font = Font(color="FFFF00")

    for row in (5, 8, 11, 15, 18, 21, 27, 28, 35, 36):
        for column in range(3, 16, 2):
            worksheet.cell(row, column, "XXXXXXXXXXXX")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _records() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "employee_name": "AM Med Tech",
                "role": "Med Tech",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Assisted Living",
                "floor": "AL1",
                "validation_status": "Valid",
            },
            {
                "employee_name": "PM Med Tech",
                "role": "Med Tech",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Assisted Living",
                "floor": "AL1",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Caregiver A",
                "role": "Caregiver",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Assisted Living",
                "floor": "AL1",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Caregiver B",
                "role": "Caregiver",
                "shift_date": "2026-07-31",
                "start_time": "07:00",
                "unit": "Assisted Living",
                "floor": "AL1",
                "validation_status": "Valid",
            },
            {
                "employee_name": "OPEN",
                "role": "Caregiver",
                "shift_date": "2026-07-31",
                "start_time": "19:00",
                "unit": "Assisted Living",
                "floor": "AL1",
                "validation_status": "Valid",
            },
            {
                "employee_name": "Week Two",
                "role": "Caregiver",
                "shift_date": "2026-08-07",
                "start_time": "07:00",
                "unit": "Memory Care",
                "floor": "MC1",
                "validation_status": "Valid",
            },
        ]
    )


def test_populates_dates_slots_and_preserves_blocked_cells() -> None:
    output = export_al_mc_template(_records(), _template_bytes())
    worksheet = load_workbook(BytesIO(output), data_only=False).active

    assert worksheet["B2"].value == "Friday, 7/31 - AM"
    assert worksheet["O12"].value == "Thursday, 8/13 - PM"
    assert worksheet["B3"].value == "AM Med Tech"
    assert worksheet["C3"].value == "PM Med Tech"
    assert worksheet["B4"].value == "Caregiver A"
    assert worksheet["B5"].value == "Caregiver B"
    assert worksheet["C4"].value == "OPEN"
    assert worksheet["C4"].font.bold is True
    assert worksheet["C4"].font.color.rgb == "FFFF0000"
    assert worksheet["C5"].value == "XXXXXXXXXXXX"
    assert worksheet["C5"].font.color.rgb == "FFFFFF00"
    assert worksheet["B6"].value == "OPEN"
    assert worksheet["B6"].font.bold is True
    assert worksheet["B6"].font.color.rgb == "FFFF0000"
    assert worksheet["B3"].font.bold is False
    assert worksheet["B3"].font.color.rgb == "FF000000"
    assert worksheet["C3"].font.bold is False
    assert worksheet["C3"].font.color.rgb == "FFFFFFFF"
    assert worksheet["B34"].value == "Week Two"
    assert worksheet["C3"].fill.fgColor.rgb == "00000000"
    assert worksheet["B5"].alignment.wrap_text is True
    assert worksheet.row_dimensions[5].height == 48
    assert worksheet.column_dimensions["A"].width == 14
    assert worksheet.column_dimensions["B"].width == 22
    assert worksheet["A3"].font.sz == 12
    assert worksheet["B2"].font.sz == 10
    assert worksheet["B1"].font.sz == 24


def test_wraps_slot_overflow_without_dropping_names() -> None:
    records = pd.concat(
        [
            _records(),
            pd.DataFrame(
                [
                    {
                        "employee_name": "Extra PM",
                        "role": "Caregiver",
                        "shift_date": "2026-07-31",
                        "start_time": "18:00",
                        "unit": "Assisted Living",
                        "floor": "AL1",
                        "validation_status": "Valid",
                    }
                ]
            ),
        ],
        ignore_index=True,
    )

    output = export_al_mc_template(records, _template_bytes())
    worksheet = load_workbook(BytesIO(output)).active

    assert worksheet["C4"].value == "Extra PM\nOPEN"
    assert worksheet["C4"].alignment.wrap_text is True
    assert worksheet["C4"].font.bold is False
    assert worksheet["C4"].font.color.rgb == "FFFFFFFF"
    assert worksheet.row_dimensions[4].height == 48
    assert worksheet["C5"].value == "XXXXXXXXXXXX"


def test_excludes_invalid_records_and_rejects_wrong_template() -> None:
    records = _records()
    records.loc[0, "validation_status"] = "Invalid"
    worksheet = load_workbook(
        BytesIO(export_al_mc_template(records, _template_bytes()))
    ).active
    assert worksheet["B3"].value == "OPEN"
    assert worksheet["C3"].value == "PM Med Tech"

    records["validation_status"] = "Invalid"
    with pytest.raises(ALMCTemplateError, match="no valid reviewed records"):
        export_al_mc_template(records, _template_bytes())

    blank = Workbook()
    payload = BytesIO()
    blank.save(payload)
    with pytest.raises(ALMCTemplateError, match="does not match"):
        export_al_mc_template(_records(), payload.getvalue())


def test_template_filename_is_windows_safe() -> None:
    assert al_mc_template_filename("AL MC report.pdf") == (
        "AL_MC_report_AL_MC_template.xlsx"
    )


def test_built_in_template_is_available() -> None:
    assert BUILT_IN_TEMPLATE_PATH.is_file()
    output = export_al_mc_template(_records())
    assert load_workbook(BytesIO(output)).active["B3"].value == "AM Med Tech"
