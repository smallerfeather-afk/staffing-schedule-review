"""Generate cleaned Excel workbooks from reviewed schedule records."""

from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from staffing_review.constants import (
    CANONICAL_FIELDS,
    PDF_METADATA_COLUMNS,
    REVIEW_COLUMNS,
)
from staffing_review.services.mapping import (
    MissingColumnsError,
    resolve_column_mapping,
)

SHEET_NAME = "Reviewed Schedule"
HEADER_FILL = PatternFill("solid", fgColor="183A35")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
HEADER_BORDER = Border(bottom=Side(style="medium", color="D99B45"))
INVALID_FILL = PatternFill("solid", fgColor="FCE8E6")
INVALID_FONT = Font(name="Aptos", size=10, color="9C0006")
OPEN_FILL = PatternFill("solid", fgColor="FFF1CC")
OPEN_FONT = Font(name="Aptos", size=10, color="7A4B00", bold=True)
BODY_FONT = Font(name="Aptos", size=10, color="1F2933")

DEFAULT_WIDTHS = {
    "source_row": 12,
    "employee_name": 24,
    "role": 18,
    "shift_date": 14,
    "start_time": 12,
    "unit": 22,
    "floor": 14,
    "validation_status": 18,
    "validation_messages": 56,
    "source_page": 12,
    "extraction_confidence": 20,
    "extraction_notes": 48,
    "extraction_reviewed": 20,
}


def _copy_cell_style(source: Cell, target: Cell) -> None:
    if not source.has_style:
        return
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def _safe_text(value: object) -> str:
    """Prevent edited text from becoming an Excel formula."""
    text = "" if pd.isna(value) else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _excel_value(field: str, value: object) -> object:
    if pd.isna(value):
        return None
    text = str(value)

    if field in {"source_row", "source_page"}:
        try:
            return int(float(text))
        except ValueError:
            return _safe_text(value)

    if field == "extraction_confidence":
        try:
            return float(value)
        except (TypeError, ValueError):
            return _safe_text(value)

    if field == "extraction_reviewed":
        return bool(value)

    if field == "shift_date":
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            return _safe_text(value)

    if field == "start_time":
        try:
            return datetime.strptime(text, "%H:%M").time()
        except ValueError:
            return _safe_text(value)

    return _safe_text(value)


def _is_open_shift(employee_name: object) -> bool:
    name = "" if pd.isna(employee_name) else str(employee_name).strip().casefold()
    return not name or name == "open" or name.startswith("open ")


def _source_styles(
    source_workbook: bytes | None,
) -> tuple[object | None, dict[str, int]]:
    if not source_workbook:
        return None, {}

    try:
        workbook = load_workbook(
            BytesIO(source_workbook),
            data_only=True,
            keep_links=False,
        )
        worksheet = workbook.worksheets[0]
        headers = pd.Index(
            worksheet.cell(row=1, column=column).value
            for column in range(1, worksheet.max_column + 1)
        )
        header_mapping = resolve_column_mapping(headers)
    except (KeyError, MissingColumnsError, OSError, ValueError):
        return None, {}

    source_columns: dict[str, int] = {}
    for column, header in enumerate(headers, start=1):
        canonical = header_mapping.get(header)
        if canonical and canonical not in source_columns:
            source_columns[canonical] = column

    return worksheet, source_columns


def _apply_default_layout(
    worksheet: object,
    row_count: int,
    export_columns: list[str],
) -> None:
    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(len(export_columns))
    worksheet.auto_filter.ref = (
        f"A1:{last_column}{max(row_count + 1, 1)}"
    )
    worksheet.sheet_view.showGridLines = False
    worksheet.print_title_rows = "1:1"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    for column_index, field in enumerate(export_columns, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = (
            DEFAULT_WIDTHS[field]
        )

    worksheet.row_dimensions[1].height = 25


def _apply_source_formatting(
    output_sheet: object,
    source_sheet: object | None,
    source_columns: dict[str, int],
    records: pd.DataFrame,
    export_columns: list[str],
) -> None:
    if source_sheet is None:
        return

    for field, source_column in source_columns.items():
        output_column = export_columns.index(field) + 1
        _copy_cell_style(
            source_sheet.cell(row=1, column=source_column),
            output_sheet.cell(row=1, column=output_column),
        )
        source_width = source_sheet.column_dimensions[
            get_column_letter(source_column)
        ].width
        if source_width:
            output_sheet.column_dimensions[
                get_column_letter(output_column)
            ].width = min(max(source_width, 10), 56)

    for output_row, (_, record) in enumerate(records.iterrows(), start=2):
        try:
            source_row = int(record["source_row"])
        except (KeyError, TypeError, ValueError):
            continue

        source_height = source_sheet.row_dimensions[source_row].height
        if source_height:
            output_sheet.row_dimensions[output_row].height = source_height

        for field, source_column in source_columns.items():
            output_column = export_columns.index(field) + 1
            _copy_cell_style(
                source_sheet.cell(row=source_row, column=source_column),
                output_sheet.cell(row=output_row, column=output_column),
            )


def export_reviewed_schedule(
    records: pd.DataFrame,
    source_workbook: bytes | None = None,
) -> bytes:
    """Return a formatted `.xlsx` workbook containing current reviewed data."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    export_columns = list(REVIEW_COLUMNS) + [
        column
        for column in PDF_METADATA_COLUMNS
        if column in records.columns
    ]

    for column, field in enumerate(export_columns, start=1):
        cell = worksheet.cell(row=1, column=column, value=field)
        cell.fill = copy(HEADER_FILL)
        cell.font = copy(HEADER_FONT)
        cell.border = copy(HEADER_BORDER)
        cell.alignment = Alignment(vertical="center")

    for row_number, (_, record) in enumerate(records.iterrows(), start=2):
        is_open = _is_open_shift(record.get("employee_name"))
        status = str(record.get("validation_status", ""))

        for column, field in enumerate(export_columns, start=1):
            value = record.get(field)
            if field == "validation_status" and is_open:
                value = "Open - Invalid" if status == "Invalid" else "Open"
            cell = worksheet.cell(
                row=row_number,
                column=column,
                value=_excel_value(field, value),
            )
            cell.font = copy(BODY_FONT)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=field
                in {"validation_messages", "extraction_notes"},
            )

    _apply_default_layout(worksheet, len(records), export_columns)
    source_sheet, source_columns = _source_styles(source_workbook)
    _apply_source_formatting(
        worksheet,
        source_sheet,
        source_columns,
        records,
        export_columns,
    )

    date_column = export_columns.index("shift_date") + 1
    time_column = export_columns.index("start_time") + 1
    for row_number, (_, record) in enumerate(records.iterrows(), start=2):
        worksheet.cell(row=row_number, column=date_column).number_format = "yyyy-mm-dd"
        worksheet.cell(row=row_number, column=time_column).number_format = "hh:mm"

        is_open = _is_open_shift(record.get("employee_name"))
        is_invalid = record.get("validation_status") == "Invalid"
        if not is_open and not is_invalid:
            continue

        fill = OPEN_FILL if is_open else INVALID_FILL
        font = OPEN_FONT if is_open else INVALID_FONT
        for column in range(1, len(export_columns) + 1):
            cell = worksheet.cell(row=row_number, column=column)
            cell.fill = copy(fill)
            cell.font = copy(font)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_filename(source_label: str) -> str:
    """Create a Windows-safe download name from the current source."""
    stem = Path(source_label).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._")
    return f"{cleaned or 'staffing_schedule'}_reviewed.xlsx"
