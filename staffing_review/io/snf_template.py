"""Populate the approved two-week SNF schedule template."""

from collections import defaultdict
from copy import copy
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
import re

from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

BUILT_IN_TEMPLATE_PATH = (
    Path(__file__).parents[2]
    / "templates"
    / "SNF Empty 2 Week Template.xlsx"
)

SCHEDULE_ROW_HEIGHT = 48
LABEL_COLUMN_WIDTH = 14
SHIFT_COLUMN_WIDTH = 22
BODY_FONT_SIZE = 11
LABEL_FONT_SIZE = 12
HEADER_FONT_SIZE = 10
TITLE_FONT_SIZE = 24
AM_FONT_COLOR = "FF000000"
PM_FONT_COLOR = "FFFFFFFF"
OPEN_FONT_COLOR = "FFFF0000"
BLOCKED_FONT_COLOR = "FFFFFF00"
RN_FONT_COLOR = "FF7030A0"

CNA_ROW_LABELS = (
    "James",
    "James/Johns",
    "Johns",
    "Folly",
    "Folly/Morris",
    "Morris",
)
NURSE_ROW_LABELS = ("James", "Johns", "Folly", "Morris")
FLOOR_TO_ROW = {
    "s1james": "James",
    "s1jamesjohns": "James/Johns",
    "s1johns": "Johns",
    "s2folly": "Folly",
    "s2follymorris": "Folly/Morris",
    "s2morris": "Morris",
}


class SNFTemplateError(ValueError):
    """Raised when reviewed records cannot safely populate the SNF template."""


def _normalized_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _safe_name(value: object) -> str:
    name = "" if pd.isna(value) else str(value).strip()
    if not name or name.casefold().startswith("open"):
        return "OPEN"
    if name.startswith(("=", "+", "-", "@")):
        return f"'{name}"
    return name


def _is_blocked(value: object) -> bool:
    return str(value or "").strip().upper().startswith("XXXX")


def _set_cell_font(
    cell: object,
    *,
    color: str,
    bold: bool,
) -> None:
    font = copy(cell.font)
    font.name = "Aptos"
    font.size = BODY_FONT_SIZE
    font.bold = bold
    font.color = color
    cell.font = font


def _set_nurse_names(
    cell: object,
    entries: list[tuple[str, int, str, str]],
    *,
    period: str,
) -> None:
    """Write nurse names with purple applied only to RN text runs."""
    if all(entry[2] == "OPEN" for entry in entries):
        cell.value = "OPEN"
        _set_cell_font(cell, color=OPEN_FONT_COLOR, bold=True)
        return

    default_color = AM_FONT_COLOR if period == "AM" else PM_FONT_COLOR
    rich_text = CellRichText()
    for index, (_, _, name, role) in enumerate(entries):
        if index:
            rich_text.append("\n")
        rich_text.append(
            TextBlock(
                InlineFont(
                    rFont="Aptos",
                    sz=BODY_FONT_SIZE,
                    b=False,
                    color=RN_FONT_COLOR if role == "RN" else default_color,
                ),
                name,
            )
        )
    cell.value = rich_text
    _set_cell_font(cell, color=default_color, bold=False)


def _template_layout(
    worksheet: object,
) -> tuple[
    dict[tuple[str, int, str], int],
    dict[tuple[str, int], int],
]:
    row_map: dict[tuple[str, int, str], int] = {}
    expected_sections = (
        ("CNA", 0, 2, CNA_ROW_LABELS),
        ("CNA", 1, 9, CNA_ROW_LABELS),
        ("NURSE", 0, 17, NURSE_ROW_LABELS),
        ("NURSE", 1, 22, NURSE_ROW_LABELS),
    )
    header_rows: dict[tuple[str, int], int] = {}
    for section, week, header_row, labels in expected_sections:
        if (
            str(worksheet.cell(header_row, 1).value or "").strip().casefold()
            != "hall"
        ):
            raise SNFTemplateError(
                "This workbook does not match the approved SNF two-week "
                "blank template layout."
            )
        header_rows[(section, week)] = header_row
        for offset, label in enumerate(labels, start=1):
            actual = str(
                worksheet.cell(header_row + offset, 1).value or ""
            ).strip()
            if _normalized_key(actual) != _normalized_key(label):
                raise SNFTemplateError(
                    "This workbook does not match the approved SNF two-week "
                    "blank template layout."
                )
            row_map[(section, week, label)] = header_row + offset
    return row_map, header_rows


def _format_grid(
    worksheet: object,
    row_map: dict[tuple[str, int, str], int],
    header_rows: dict[tuple[str, int], int],
) -> None:
    worksheet.column_dimensions["A"].width = LABEL_COLUMN_WIDTH
    for column in range(2, 16):
        worksheet.column_dimensions[get_column_letter(column)].width = (
            SHIFT_COLUMN_WIDTH
        )

    for row_number in sorted(set(row_map.values())):
        worksheet.row_dimensions[row_number].height = SCHEDULE_ROW_HEIGHT
        label_cell = worksheet.cell(row_number, 1)
        label_font = copy(label_cell.font)
        label_font.name = "Aptos"
        label_font.size = LABEL_FONT_SIZE
        label_font.bold = True
        label_font.color = AM_FONT_COLOR
        label_cell.font = label_font
        label_alignment = copy(label_cell.alignment)
        label_alignment.horizontal = "left"
        label_alignment.vertical = "center"
        label_cell.alignment = label_alignment

        for column in range(2, 16):
            cell = worksheet.cell(row_number, column)
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = True
            cell.alignment = alignment
            if _is_blocked(cell.value):
                _set_cell_font(
                    cell,
                    color=BLOCKED_FONT_COLOR,
                    bold=False,
                )
            else:
                _set_cell_font(
                    cell,
                    color=AM_FONT_COLOR if column % 2 == 0 else PM_FONT_COLOR,
                    bold=False,
                )

    for header_row in header_rows.values():
        worksheet.row_dimensions[header_row].height = 30
        for column in range(1, 16):
            cell = worksheet.cell(header_row, column)
            font = copy(cell.font)
            font.name = "Aptos"
            font.size = HEADER_FONT_SIZE
            font.bold = True
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = False
            cell.alignment = alignment

    for title_row in (1, 16):
        worksheet.row_dimensions[title_row].height = 46
        title_cell = worksheet.cell(title_row, 2)
        font = copy(title_cell.font)
        font.name = "Aptos"
        font.size = TITLE_FONT_SIZE
        font.bold = True
        title_cell.font = font

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "B2"
    worksheet.print_area = "A1:O26"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_options.horizontalCentered = True


def _load_template() -> bytes:
    try:
        return BUILT_IN_TEMPLATE_PATH.read_bytes()
    except OSError as exc:
        raise SNFTemplateError(
            "The built-in SNF template is missing from the application."
        ) from exc


def _parse_records(
    records: pd.DataFrame,
) -> tuple[
    datetime,
    dict[tuple[str, int, str, str], list[tuple[str, int, str, str]]],
]:
    if records.empty:
        raise SNFTemplateError("There are no reviewed records to export.")
    statuses = records.get("validation_status")
    if statuses is not None:
        records = records.loc[
            statuses.astype(str).str.strip().str.casefold() != "invalid"
        ].copy()
    if records.empty:
        raise SNFTemplateError(
            "There are no valid reviewed records to export to the SNF "
            "template."
        )

    parsed: list[dict[str, object]] = []
    for _, record in records.iterrows():
        role = str(record.get("role") or "").strip().upper()
        floor_value = str(record.get("floor") or "").strip()
        floor_label = FLOOR_TO_ROW.get(_normalized_key(floor_value))
        is_house_rn = role == "RN" and not floor_value

        if role not in {"CNA", "LPN", "RN"}:
            raise SNFTemplateError(
                f"Role {role or '(blank)'} is not supported by the SNF "
                "template."
            )
        if role == "CNA" and not floor_label:
            raise SNFTemplateError(
                f"Floor {floor_value or '(blank)'} is not supported by the "
                "SNF CNA template."
            )
        if role in {"LPN", "RN"} and not floor_label and not is_house_rn:
            raise SNFTemplateError(
                f"Floor {floor_value or '(blank)'} is not supported by the "
                "SNF Nurses template."
            )

        try:
            shift_date = datetime.strptime(
                str(record.get("shift_date")), "%Y-%m-%d"
            )
            start_time = datetime.strptime(
                str(record.get("start_time")), "%H:%M"
            )
        except ValueError as exc:
            raise SNFTemplateError(
                "All shift dates and start times must be normalized before "
                "filling the SNF template."
            ) from exc

        name = _safe_name(record.get("employee_name"))
        parsed.append(
            {
                "date": shift_date,
                "time": start_time.strftime("%H:%M"),
                "period": "AM" if start_time.hour < 12 else "PM",
                "role": role,
                "floor_label": floor_label,
                "is_house_rn": is_house_rn,
                "name": name,
            }
        )

    start_date = min(item["date"] for item in parsed)
    groups: defaultdict[
        tuple[str, int, str, str],
        list[tuple[str, int, str, str]],
    ] = defaultdict(list)
    for item in parsed:
        day_offset = (item["date"] - start_date).days
        if day_offset < 0 or day_offset > 13:
            raise SNFTemplateError(
                "The SNF template supports one continuous 14-day period."
            )
        role = str(item["role"])
        period = str(item["period"])
        time_value = str(item["time"])
        name = str(item["name"])

        if role == "CNA":
            labels = (str(item["floor_label"]),)
            section = "CNA"
            display_name = name
            role_rank = 0
        elif item["is_house_rn"]:
            labels = NURSE_ROW_LABELS
            section = "NURSE"
            display_name = name
            role_rank = 2
        else:
            labels = (str(item["floor_label"]),)
            section = "NURSE"
            display_name = name
            role_rank = 1 if role == "RN" else 0

        for label in labels:
            groups[(section, day_offset, label, period)].append(
                (time_value, role_rank, display_name, role)
            )

    for entries in groups.values():
        entries.sort(key=lambda entry: (entry[0], entry[1], entry[2].casefold()))
    return start_date, groups


def export_snf_template(records: pd.DataFrame) -> bytes:
    """Return a populated copy of the approved SNF two-week template."""
    try:
        workbook = load_workbook(BytesIO(_load_template()))
    except SNFTemplateError:
        raise
    except Exception as exc:
        raise SNFTemplateError(
            "The built-in SNF template workbook could not be opened."
        ) from exc

    worksheet = workbook.worksheets[0]
    row_map, header_rows = _template_layout(worksheet)
    start_date, groups = _parse_records(records)
    _format_grid(worksheet, row_map, header_rows)

    for (_, week), header_row in header_rows.items():
        for day in range(7):
            shift_date = start_date + timedelta(days=(week * 7) + day)
            for period_offset, period in enumerate(("AM", "PM")):
                column = 2 + (day * 2) + period_offset
                worksheet.cell(
                    row=header_row,
                    column=column,
                    value=(
                        f"{shift_date:%A}, {shift_date.month}/"
                        f"{shift_date.day} - {period}"
                    ),
                )

    for row_number in row_map.values():
        for column in range(2, 16):
            cell = worksheet.cell(row_number, column)
            if not _is_blocked(cell.value):
                cell.value = None

    folly_morris_merges: dict[
        tuple[int, int], list[tuple[str, int, str, str]]
    ] = {}
    for day_offset in range(14):
        folly_key = ("NURSE", day_offset, "Folly", "PM")
        morris_key = ("NURSE", day_offset, "Morris", "PM")
        folly_entries = groups.get(folly_key, [])
        morris_names = [
            entry for entry in groups.get(morris_key, [])
            if entry[2] != "OPEN"
        ]
        folly_is_open = not folly_entries or all(
            entry[2] == "OPEN" for entry in folly_entries
        )
        if folly_is_open and morris_names:
            folly_morris_merges[(day_offset // 7, day_offset % 7)] = (
                morris_names
            )

    for (section, day_offset, label, period), entries in groups.items():
        week = day_offset // 7
        day = day_offset % 7
        if (
            section == "NURSE"
            and period == "PM"
            and label in {"Folly", "Morris"}
            and (week, day) in folly_morris_merges
        ):
            continue
        row_number = row_map[(section, week, label)]
        column = 2 + (day * 2) + (1 if period == "PM" else 0)
        cell = worksheet.cell(row_number, column)
        if _is_blocked(cell.value):
            date_label = (start_date + timedelta(days=day_offset)).strftime(
                "%Y-%m-%d"
            )
            raise SNFTemplateError(
                f"{label} has a {section} {period} shift on {date_label}, "
                "but that template cell is blocked."
            )
        if section == "NURSE":
            _set_nurse_names(cell, entries, period=period)
        else:
            names = [entry[2] for entry in entries]
            cell.value = "\n".join(names)
            if all(name == "OPEN" for name in names):
                _set_cell_font(cell, color=OPEN_FONT_COLOR, bold=True)
            else:
                _set_cell_font(
                    cell,
                    color=AM_FONT_COLOR if period == "AM" else PM_FONT_COLOR,
                    bold=False,
                )

    for (week, day), entries in folly_morris_merges.items():
        column = 3 + (day * 2)
        folly_row = row_map[("NURSE", week, "Folly")]
        morris_row = row_map[("NURSE", week, "Morris")]
        cell = worksheet.cell(folly_row, column)
        _set_nurse_names(cell, entries, period="PM")
        worksheet.merge_cells(
            start_row=folly_row,
            start_column=column,
            end_row=morris_row,
            end_column=column,
        )

    for row_number in row_map.values():
        for column in range(2, 16):
            cell = worksheet.cell(row_number, column)
            if isinstance(cell, MergedCell):
                continue
            if not _is_blocked(cell.value) and not cell.value:
                cell.value = "OPEN"
                _set_cell_font(cell, color=OPEN_FONT_COLOR, bold=True)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def snf_template_filename(source_label: str) -> str:
    """Create a Windows-safe populated-template download name."""
    stem = Path(source_label).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._")
    return f"{cleaned or 'snf_schedule'}_SNF_template.xlsx"
