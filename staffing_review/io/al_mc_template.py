"""Populate the approved two-week AL/MC schedule template."""

from collections import defaultdict
from copy import copy
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

ROW_LABEL_PATTERN = re.compile(
    r"^(?P<floor>AL[1-3]|MC[1-2])\s*-\s*"
    r"(?P<role>MT|CG)(?P<slot>\d*)\s*$",
    re.IGNORECASE,
)
ROLE_CODES = {
    "caregiver": "CG",
    "cg": "CG",
    "med tech": "MT",
    "medication technician": "MT",
    "mt": "MT",
}
BUILT_IN_TEMPLATE_PATH = (
    Path(__file__).parents[2]
    / "templates"
    / "AL-MC 2 Week Blank Template.xlsx"
)
SCHEDULE_ROW_HEIGHT = 48
LABEL_COLUMN_WIDTH = 14
SHIFT_COLUMN_WIDTH = 22
BODY_FONT_SIZE = 11
LABEL_FONT_SIZE = 12
HEADER_FONT_SIZE = 10
TITLE_FONT_SIZE = 24
AM_FONT_COLOR = "FF000000"
PM_FONT_COLOR = "FFFFFFFF"
OPEN_FONT_COLOR = "FFFF0000"
BLOCKED_FONT_COLOR = "FFFFFF00"


class ALMCTemplateError(ValueError):
    """Raised when reviewed records cannot safely populate the AL/MC template."""


def _safe_name(value: object) -> str:
    name = "" if pd.isna(value) else str(value).strip()
    if not name or name.casefold().startswith("open"):
        return "OPEN"
    if name.startswith(("=", "+", "-", "@")):
        return f"'{name}"
    return name


def _template_layout(worksheet: object) -> tuple[
    dict[tuple[int, str, str, int], int],
    dict[tuple[str, int], int],
]:
    rows: dict[tuple[int, str, str, int], int] = {}
    occurrences: defaultdict[tuple[str, str, int], int] = defaultdict(int)

    for row_number in range(1, worksheet.max_row + 1):
        match = ROW_LABEL_PATTERN.fullmatch(
            str(worksheet.cell(row=row_number, column=1).value or "").strip()
        )
        if not match:
            continue
        floor = match.group("floor").upper()
        role = match.group("role").upper()
        slot = int(match.group("slot") or "1")
        occurrence_key = (floor, role, slot)
        week = occurrences[occurrence_key]
        occurrences[occurrence_key] += 1
        rows[(week, floor, role, slot)] = row_number

    header_rows: dict[tuple[str, int], int] = {}
    header_occurrences: defaultdict[str, int] = defaultdict(int)
    for row_number in range(1, worksheet.max_row + 1):
        if str(worksheet.cell(row=row_number, column=1).value or "").strip() != "Hall":
            continue
        next_label = str(
            worksheet.cell(row=row_number + 1, column=1).value or ""
        ).strip()
        match = ROW_LABEL_PATTERN.fullmatch(next_label)
        if not match:
            continue
        section = match.group("floor")[:2].upper()
        week = header_occurrences[section]
        header_occurrences[section] += 1
        header_rows[(section, week)] = row_number

    required_rows = {
        (week, floor, role, slot)
        for week in (0, 1)
        for floor, caregiver_slots in {
            "AL1": 2,
            "AL2": 2,
            "AL3": 2,
            "MC1": 3,
            "MC2": 2,
        }.items()
        for role, slots in (("MT", 1), ("CG", caregiver_slots))
        for slot in range(1, slots + 1)
    }
    required_headers = {
        (section, week)
        for section in ("AL", "MC")
        for week in (0, 1)
    }
    if not required_rows.issubset(rows) or not required_headers.issubset(
        header_rows
    ):
        raise ALMCTemplateError(
            "This workbook does not match the approved AL/MC two-week "
            "blank template layout."
        )
    return rows, header_rows


def _parse_records(records: pd.DataFrame) -> tuple[
    datetime,
    dict[tuple[int, str, str, str], list[tuple[str, str]]],
]:
    if records.empty:
        raise ALMCTemplateError("There are no reviewed records to export.")
    statuses = records.get("validation_status")
    if statuses is not None:
        records = records.loc[
            statuses.astype(str).str.strip().str.casefold() != "invalid"
        ].copy()
    if records.empty:
        raise ALMCTemplateError(
            "There are no valid reviewed records to export to the AL/MC "
            "template."
        )

    parsed: list[dict[str, object]] = []
    for _, record in records.iterrows():
        try:
            shift_date = datetime.strptime(
                str(record.get("shift_date")), "%Y-%m-%d"
            )
            start_time = datetime.strptime(
                str(record.get("start_time")), "%H:%M"
            )
        except ValueError as exc:
            raise ALMCTemplateError(
                "All shift dates and start times must be normalized before "
                "filling the AL/MC template."
            ) from exc

        floor = str(record.get("floor") or "").strip().upper()
        if floor not in {"AL1", "AL2", "AL3", "MC1", "MC2"}:
            raise ALMCTemplateError(
                f"Floor {floor or '(blank)'} is not supported by the "
                "AL/MC template."
            )
        role = ROLE_CODES.get(str(record.get("role") or "").strip().casefold())
        if not role:
            raise ALMCTemplateError(
                f"Role {record.get('role') or '(blank)'} is not supported by "
                "the AL/MC template."
            )
        parsed.append(
            {
                "date": shift_date,
                "time": start_time.strftime("%H:%M"),
                "period": "AM" if start_time.hour < 12 else "PM",
                "floor": floor,
                "role": role,
                "name": _safe_name(record.get("employee_name")),
            }
        )

    start_date = min(item["date"] for item in parsed)
    groups: defaultdict[
        tuple[int, str, str, str], list[tuple[str, str]]
    ] = defaultdict(list)
    for item in parsed:
        day_offset = (item["date"] - start_date).days
        if day_offset < 0 or day_offset > 13:
            raise ALMCTemplateError(
                "The AL/MC template supports one continuous 14-day period."
            )
        groups[
            (
                day_offset,
                str(item["floor"]),
                str(item["role"]),
                str(item["period"]),
            )
        ].append((str(item["time"]), str(item["name"])))

    for entries in groups.values():
        entries.sort(key=lambda entry: (entry[0], entry[1].casefold()))
    return start_date, groups


def _is_blocked(value: object) -> bool:
    return str(value or "").strip().upper().startswith("XXXX")


def _set_cell_font(
    cell: object,
    *,
    color: str,
    bold: bool,
) -> None:
    font = copy(cell.font)
    font.name = "Aptos"
    font.size = BODY_FONT_SIZE
    font.bold = bold
    font.color = color
    cell.font = font


def _format_schedule_grid(
    worksheet: object,
    row_map: dict[tuple[int, str, str, int], int],
) -> None:
    worksheet.column_dimensions["A"].width = LABEL_COLUMN_WIDTH
    for column in range(2, 16):
        worksheet.column_dimensions[get_column_letter(column)].width = (
            SHIFT_COLUMN_WIDTH
        )

    for row_number in sorted(set(row_map.values())):
        worksheet.row_dimensions[row_number].height = SCHEDULE_ROW_HEIGHT
        label_cell = worksheet.cell(row=row_number, column=1)
        label_font = copy(label_cell.font)
        label_font.name = "Aptos"
        label_font.size = LABEL_FONT_SIZE
        label_font.bold = True
        label_font.color = AM_FONT_COLOR
        label_cell.font = label_font
        label_alignment = copy(label_cell.alignment)
        label_alignment.horizontal = "left"
        label_alignment.vertical = "center"
        label_cell.alignment = label_alignment
        for column in range(2, 16):
            cell = worksheet.cell(row=row_number, column=column)
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = True
            cell.alignment = alignment
            if _is_blocked(cell.value):
                _set_cell_font(
                    cell,
                    color=BLOCKED_FONT_COLOR,
                    bold=False,
                )
            else:
                _set_cell_font(
                    cell,
                    color=AM_FONT_COLOR if column % 2 == 0 else PM_FONT_COLOR,
                    bold=False,
                )

    for header_row in (2, 12, 24, 32):
        worksheet.row_dimensions[header_row].height = 30
        for column in range(1, 16):
            cell = worksheet.cell(row=header_row, column=column)
            font = copy(cell.font)
            font.name = "Aptos"
            font.size = HEADER_FONT_SIZE
            font.bold = True
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = False
            cell.alignment = alignment
    worksheet.row_dimensions[1].height = 46
    worksheet.row_dimensions[23].height = 46
    for title_cell in (worksheet["B1"], worksheet["B23"]):
        font = copy(title_cell.font)
        font.name = "Aptos"
        font.size = TITLE_FONT_SIZE
        font.bold = True
        title_cell.font = font
    worksheet.row_dimensions[22].height = 8
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "B2"
    worksheet.print_area = "A1:O39"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_options.horizontalCentered = True
    worksheet.sheet_properties.pageSetUpPr.autoPageBreaks = False


def _load_template(template: bytes | None) -> bytes:
    if template:
        return template
    try:
        return BUILT_IN_TEMPLATE_PATH.read_bytes()
    except OSError as exc:
        raise ALMCTemplateError(
            "The built-in AL/MC template is missing from the application."
        ) from exc


def export_al_mc_template(
    records: pd.DataFrame,
    template: bytes | None = None,
) -> bytes:
    """Return a populated copy of the approved AL/MC two-week template."""
    try:
        workbook = load_workbook(BytesIO(_load_template(template)))
    except Exception as exc:
        raise ALMCTemplateError(
            "The AL/MC template workbook could not be opened."
        ) from exc

    worksheet = workbook.worksheets[0]
    row_map, header_rows = _template_layout(worksheet)
    start_date, groups = _parse_records(records)
    _format_schedule_grid(worksheet, row_map)

    for (section, week), header_row in header_rows.items():
        for day in range(7):
            shift_date = start_date + timedelta(days=(week * 7) + day)
            for period_offset, period in enumerate(("AM", "PM")):
                column = 2 + (day * 2) + period_offset
                worksheet.cell(
                    row=header_row,
                    column=column,
                    value=(
                        f"{shift_date:%A}, {shift_date.month}/"
                        f"{shift_date.day} - {period}"
                    ),
                )

    for row_number in row_map.values():
        for column in range(2, 16):
            cell = worksheet.cell(row=row_number, column=column)
            if not _is_blocked(cell.value):
                cell.value = None

    for (day_offset, floor, role, period), entries in groups.items():
        week = day_offset // 7
        day = day_offset % 7
        column = 2 + (day * 2) + (1 if period == "PM" else 0)
        candidate_rows = sorted(
            (
                slot,
                row_number,
            )
            for (row_week, row_floor, row_role, slot), row_number
            in row_map.items()
            if row_week == week
            and row_floor == floor
            and row_role == role
            and not _is_blocked(
                worksheet.cell(row=row_number, column=column).value
            )
        )
        if not candidate_rows:
            raise ALMCTemplateError(
                f"{floor} has no available {role} {period} template slot."
            )

        assignments: list[tuple[int, list[str]]] = [
            (row_number, [entries[index][1]])
            for index, (_, row_number) in enumerate(candidate_rows)
            if index < len(entries)
        ]
        if len(entries) > len(candidate_rows):
            assignments[-1][1].extend(
                name for _, name in entries[len(candidate_rows) :]
            )

        for row_number, names in assignments:
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = "\n".join(names)
            if all(name == "OPEN" for name in names):
                _set_cell_font(cell, color=OPEN_FONT_COLOR, bold=True)
            else:
                _set_cell_font(
                    cell,
                    color=(
                        AM_FONT_COLOR if period == "AM" else PM_FONT_COLOR
                    ),
                    bold=False,
                )

    for row_number in row_map.values():
        for column in range(2, 16):
            cell = worksheet.cell(row=row_number, column=column)
            if not _is_blocked(cell.value) and not cell.value:
                cell.value = "OPEN"
                _set_cell_font(cell, color=OPEN_FONT_COLOR, bold=True)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def al_mc_template_filename(source_label: str) -> str:
    """Create a Windows-safe populated-template download name."""
    stem = Path(source_label).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._")
    return f"{cleaned or 'al_mc_schedule'}_AL_MC_template.xlsx"
